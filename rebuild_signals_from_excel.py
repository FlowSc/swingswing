# ============================================================
#  Rebuild swing signal JSON from the latest swing Excel output.
#  Use this when FinanceDataReader/KRX network access fails.
# ============================================================

from __future__ import annotations

from datetime import datetime
import json
from pathlib import Path

from openpyxl import load_workbook


OUTPUT_DIR = Path("output")
TOP_N = 30


def latest_excel_file() -> Path:
    files = sorted(OUTPUT_DIR.glob("kospi_swing_*.xlsx"))
    if not files:
        raise FileNotFoundError("No output/kospi_swing_*.xlsx file found.")
    return files[-1]


def main() -> None:
    OUTPUT_DIR.mkdir(exist_ok=True)
    excel_file = latest_excel_file()
    workbook = load_workbook(excel_file, read_only=True, data_only=True)
    if "TOP30" not in workbook.sheetnames:
        raise ValueError(f"Sheet TOP30 not found in {excel_file}")

    sheet = workbook["TOP30"]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value) for value in next(rows)]
    signals = []
    for row in rows:
        item = {
            header: value
            for header, value in zip(headers, row)
            if header and value is not None
        }
        if item:
            signals.append(item)
        if len(signals) >= TOP_N:
            break

    today = datetime.today()
    payload = {
        "date": today.strftime("%Y-%m-%d"),
        "strategy": "swing_ichimoku_bollinger",
        "top_n": TOP_N,
        "source": str(excel_file),
        "signals": signals,
    }
    output_file = OUTPUT_DIR / f"swing_signals_{today.strftime('%Y%m%d')}.json"
    with output_file.open("w", encoding="utf-8") as file:
        json.dump(payload, file, ensure_ascii=False, indent=2)

    print(f"Rebuilt signals from: {excel_file}")
    print(f"Saved: {output_file}")
    print(f"Signals: {len(signals)}")


if __name__ == "__main__":
    main()
